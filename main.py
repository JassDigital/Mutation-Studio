import sys, json, csv, re
from pathlib import Path
from datetime import datetime
from PySide6.QtCore import Qt, QSortFilterProxyModel
from PySide6.QtGui import QStandardItemModel, QStandardItem, QColor, QFont
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
    QFrame, QFileDialog, QMessageBox, QLineEdit, QComboBox, QTableView, QHeaderView,
    QDialog, QTextEdit, QDialogButtonBox, QAbstractItemView, QSplitter, QToolButton
)
try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

APP_DIR = Path(__file__).resolve().parent
DATA_FILE = APP_DIR / "data" / "sample_report.json"
FIELDS = ["S.No.", "Mutation No", "Mutation Date", "Mutation Status", "Registry Number", "Registry Date"]
BG = "#f4f7fb"; INK = "#17243b"; MUTED = "#6b7890"; BLUE = "#356df3"; LINE = "#e4eaf3"

def clean(v):
    return "" if v is None else str(v).strip()

def normalize_date(value):
    value = clean(value)
    if not value or value == "-": return value
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%m/%d/%Y", "%d-%m-%Y"):
        try: return datetime.strptime(value, fmt).strftime("%d/%m/%Y")
        except ValueError: pass
    return value

def parse_pasted(text):
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if not lines: return []
    delimiter = "\t" if any("\t" in ln for ln in lines) else ("," if any("," in ln for ln in lines) else None)
    if delimiter:
        import csv, io
        matrix = list(csv.reader(io.StringIO("\n".join(lines)), delimiter=delimiter))
    else:
        matrix = [re.split(r"\s{2,}", ln.strip()) for ln in lines]
    matrix = [[clean(c) for c in row] for row in matrix if row]
    aliases = {
      "s.no.":"S.No.", "s.no":"S.No.", "serial no":"S.No.", "serial number":"S.No.",
      "mutation no":"Mutation No", "mutation number":"Mutation No",
      "mutation date":"Mutation Date", "mutation status":"Mutation Status",
      "registry number":"Registry Number", "registry no":"Registry Number",
      "registry date":"Registry Date"
    }
    def canon(x): return aliases.get(clean(x).lower(), clean(x))
    first = [canon(x) for x in matrix[0]]
    has_header = any(x in FIELDS for x in first)
    if has_header:
        mapping = {name:i for i,name in enumerate(first)}
        data = matrix[1:]
    else:
        mapping = {name:i for i,name in enumerate(FIELDS)}
        data = matrix
    out = []
    for n, row in enumerate(data, 1):
        rec = {}
        for field in FIELDS:
            idx = mapping.get(field)
            rec[field] = clean(row[idx]) if idx is not None and idx < len(row) else ""
        if not any(rec.values()): continue
        rec["S.No."] = rec["S.No."] or str(n)
        for key in ("Mutation Date", "Registry Date"): rec[key] = normalize_date(rec[key])
        out.append(rec)
    return out

class PasteDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Paste report data")
        self.resize(760, 440)
        layout = QVBoxLayout(self); layout.setSpacing(12)
        title = QLabel("Paste rows from Excel")
        title.setStyleSheet("font-size:20px;font-weight:700;color:#17243b")
        hint = QLabel("Copy cells in Excel and paste below. Tab-separated columns are detected automatically. Include headers if available.")
        hint.setWordWrap(True); hint.setStyleSheet("color:#6b7890")
        self.editor = QTextEdit()
        self.editor.setPlaceholderText("S.No.\\tMutation No\\tMutation Date\\tMutation Status\\tRegistry Number\\tRegistry Date")
        self.editor.setStyleSheet("QTextEdit{background:#fbfcff;border:1px solid #dbe3ef;border-radius:10px;padding:10px;font-size:13px}")
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Cancel | QDialogButtonBox.StandardButton.Ok)
        buttons.button(QDialogButtonBox.StandardButton.Ok).setText("Import pasted rows")
        buttons.accepted.connect(self.accept); buttons.rejected.connect(self.reject)
        layout.addWidget(title); layout.addWidget(hint); layout.addWidget(self.editor,1); layout.addWidget(buttons)
    def get_text(self): return self.editor.toPlainText()

class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Mutation Pendency Dashboard")
        self.resize(1400, 850)
        self.records = []
        self.department = "Department of Revenue & Disaster Management, Haryana"
        self.generated_on = ""
        self._build_ui()
        self.load_seed()
        self.refresh()

    def _build_ui(self):
        self.setStyleSheet(f"""
          QMainWindow,QWidget#root{{background:{BG};color:{INK};font-family:'Segoe UI',Arial}}
          QLabel{{color:{INK}}}
          QPushButton{{background:#fff;border:1px solid #dfe6f0;border-radius:9px;padding:10px 14px;font-weight:600;color:#263650}}
          QPushButton:hover{{border-color:{BLUE};background:#f6f8ff}}
          QPushButton#primary{{background:{BLUE};color:white;border:0}}
          QPushButton#primary:hover{{background:#2457d6}}
          QLineEdit,QComboBox{{background:white;border:1px solid #dfe6f0;border-radius:9px;padding:10px 12px;min-height:20px;color:{INK}}}
          QTableView{{background:white;border:1px solid {LINE};border-radius:12px;gridline-color:#edf1f7;selection-background-color:#e8efff;selection-color:{INK};alternate-background-color:#fbfcfe}}
          QHeaderView::section{{background:#f7f9fd;color:#64728a;border:0;border-bottom:1px solid {LINE};padding:13px 10px;font-weight:700}}
          QComboBox QAbstractItemView{{background:white;selection-background-color:#e8efff}}
        """)
        root = QWidget(); root.setObjectName("root"); self.setCentralWidget(root)
        page = QVBoxLayout(root); page.setContentsMargins(28,24,28,22); page.setSpacing(18)
        top = QHBoxLayout()
        brand = QVBoxLayout(); brand.setSpacing(3)
        eyebrow = QLabel("HARYANA  /  REVENUE & DISASTER MANAGEMENT")
        eyebrow.setStyleSheet("font-size:10px;font-weight:800;letter-spacing:1.6px;color:#72819a")
        title = QLabel("Mutation Pendency")
        title.setStyleSheet("font-size:30px;font-weight:800;letter-spacing:-.7px")
        subtitle = QLabel("A clear, live view of mutation progress and registry details")
        subtitle.setStyleSheet("font-size:13px;color:#718098")
        brand.addWidget(eyebrow); brand.addWidget(title); brand.addWidget(subtitle)
        top.addLayout(brand); top.addStretch()
        self.import_btn = QPushButton("＋  Import Excel"); self.import_btn.setObjectName("primary"); self.import_btn.clicked.connect(self.import_excel)
        paste_btn = QPushButton("▤  Paste data"); paste_btn.clicked.connect(self.paste_data)
        export_btn = QPushButton("↓  Export CSV"); export_btn.clicked.connect(self.export_csv)
        top.addWidget(paste_btn); top.addWidget(self.import_btn); top.addWidget(export_btn)
        page.addLayout(top)
        self.source_label = QLabel("Loaded report")
        self.source_label.setStyleSheet("font-size:11px;color:#8a96a9")
        page.addWidget(self.source_label)
        cards = QHBoxLayout(); cards.setSpacing(14)
        self.kpi_total = self.make_card(cards, "TOTAL RECORDS", "—", "All mutations in this report", "▦", "#eaf0ff", "#356df3")
        self.kpi_incorp = self.make_card(cards, "INCORPORATED", "—", "Completed at incorporation stage", "✓", "#e5f8ef", "#14875c")
        self.kpi_entry = self.make_card(cards, "ENTRY PENDING", "—", "Currently at entry level", "◷", "#fff4df", "#b77812")
        self.kpi_registry = self.make_card(cards, "NO REGISTRY DETAILS", "—", "Registry number or date missing", "!", "#fff0ed", "#d45b49")
        page.addLayout(cards)
        panel = QFrame(); panel.setStyleSheet("QFrame{background:white;border:1px solid #e5ebf4;border-radius:14px}")
        panel_l = QVBoxLayout(panel); panel_l.setContentsMargins(18,16,18,14); panel_l.setSpacing(13)
        toolbar = QHBoxLayout()
        section = QVBoxLayout(); section.setSpacing(2)
        h = QLabel("Mutation register"); h.setStyleSheet("font-size:18px;font-weight:750")
        self.result_count = QLabel("Showing all records"); self.result_count.setStyleSheet("font-size:11px;color:#8290a5")
        section.addWidget(h); section.addWidget(self.result_count)
        toolbar.addLayout(section); toolbar.addStretch()
        self.search = QLineEdit(); self.search.setPlaceholderText("⌕  Search mutation, registry, status…"); self.search.setFixedWidth(300); self.search.textChanged.connect(self.apply_filters)
        self.status_filter = QComboBox(); self.status_filter.addItem("All statuses"); self.status_filter.setMinimumWidth(165); self.status_filter.currentTextChanged.connect(self.apply_filters)
        self.sort_filter = QComboBox(); self.sort_filter.addItems(["Latest mutation first", "Oldest mutation first", "Mutation number ↓", "Mutation number ↑"]); self.sort_filter.currentTextChanged.connect(self.apply_filters)
        toolbar.addWidget(self.search); toolbar.addWidget(self.status_filter); toolbar.addWidget(self.sort_filter)
        panel_l.addLayout(toolbar)
        self.model = QStandardItemModel(0, len(FIELDS)); self.model.setHorizontalHeaderLabels(FIELDS)
        self.proxy = QSortFilterProxyModel(self); self.proxy.setSourceModel(self.model); self.proxy.setFilterCaseSensitivity(Qt.CaseSensitivity.CaseInsensitive); self.proxy.setFilterKeyColumn(-1)
        self.table = QTableView(); self.table.setModel(self.proxy); self.table.setAlternatingRowColors(True)
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows); self.table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers); self.table.setSortingEnabled(False)
        self.table.verticalHeader().setVisible(False); self.table.verticalHeader().setDefaultSectionSize(42)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.table.horizontalHeader().setSectionResizeMode(0,QHeaderView.ResizeMode.ResizeToContents)
        self.table.horizontalHeader().setSectionResizeMode(1,QHeaderView.ResizeMode.ResizeToContents)
        self.table.horizontalHeader().setSectionResizeMode(3,QHeaderView.ResizeMode.ResizeToContents)
        panel_l.addWidget(self.table,1)
        bottom = QHBoxLayout()
        note = QLabel("Tip: select a status or search by mutation / registry number to narrow the list.")
        note.setStyleSheet("font-size:11px;color:#8491a5")
        self.page_info = QLabel(""); self.page_info.setStyleSheet("font-size:11px;color:#8491a5")
        bottom.addWidget(note); bottom.addStretch(); bottom.addWidget(self.page_info)
        panel_l.addLayout(bottom)
        page.addWidget(panel,1)
        foot = QHBoxLayout()
        foot.addWidget(QLabel("MUTATION PENDENCY DASHBOARD")); foot.addStretch()
        self.updated_label = QLabel(""); self.updated_label.setStyleSheet("font-size:11px;color:#8491a5")
        foot.addWidget(self.updated_label); page.addLayout(foot)
        for lab in [foot.itemAt(0).widget()]:
            lab.setStyleSheet("font-size:10px;font-weight:800;letter-spacing:1px;color:#a1adbd")

    def make_card(self, layout, label, value, desc, icon, tint, accent):
        card = QFrame(); card.setMinimumHeight(122)
        card.setStyleSheet("QFrame{background:white;border:1px solid #e5ebf4;border-radius:13px}")
        l = QVBoxLayout(card); l.setContentsMargins(16,14,16,14); l.setSpacing(5)
        row = QHBoxLayout()
        lbl = QLabel(label); lbl.setStyleSheet("font-size:10px;font-weight:800;letter-spacing:1px;color:#7c899e;border:0")
        ico = QLabel(icon); ico.setAlignment(Qt.AlignmentFlag.AlignCenter); ico.setFixedSize(30,30)
        ico.setStyleSheet(f"background:{tint};color:{accent};border-radius:9px;font-size:16px;font-weight:800;border:0")
        row.addWidget(lbl); row.addStretch(); row.addWidget(ico)
        val = QLabel(value); val.setStyleSheet(f"font-size:27px;font-weight:800;color:{INK};border:0")
        d = QLabel(desc); d.setStyleSheet("font-size:11px;color:#8491a5;border:0")
        l.addLayout(row); l.addWidget(val); l.addWidget(d); layout.addWidget(card,1)
        return val

    def load_seed(self):
        try:
            data = json.loads(DATA_FILE.read_text(encoding="utf-8"))
            self.records = data.get("records", [])
            self.department = data.get("department", self.department)
            self.generated_on = data.get("generated_on", "")
            self.source_label.setText(f"●  Sample loaded from Mutation_Pendency_Report  ·  Generated {self.generated_on}")
        except Exception as e:
            self.source_label.setText("No report loaded yet")

    def import_excel(self):
        path, _ = QFileDialog.getOpenFileName(self, "Import mutation report", "", "Excel workbooks (*.xlsx *.xlsm);;All files (*)")
        if not path: return
        if load_workbook is None:
            QMessageBox.critical(self, "Missing dependency", "Install openpyxl first: pip install openpyxl"); return
        try:
            book = load_workbook(path, data_only=True, read_only=True)
            sheet = book.active
            matrix = list(sheet.iter_rows(values_only=True))
            header_idx = None
            for i, row in enumerate(matrix[:30]):
                vals = [clean(x).lower() for x in row]
                if "mutation no" in vals and "mutation status" in vals:
                    header_idx = i; break
            if header_idx is None: raise ValueError("Could not find the column headers (Mutation No and Mutation Status).")
            raw_headers = [clean(x) for x in matrix[header_idx]]
            aliases = {"s.no.":"S.No.", "s.no":"S.No.", "mutation no":"Mutation No", "mutation date":"Mutation Date",
                       "mutation status":"Mutation Status", "registry number":"Registry Number", "registry no":"Registry Number", "registry date":"Registry Date"}
            mapping = {aliases.get(h.lower(), h):i for i,h in enumerate(raw_headers) if h}
            records=[]
            for r in matrix[header_idx+1:]:
                if not any(x is not None and clean(x) for x in r): continue
                rec={}
                for f in FIELDS:
                    idx=mapping.get(f)
                    v=r[idx] if idx is not None and idx<len(r) else ""
                    if hasattr(v,"strftime"): v=v.strftime("%d/%m/%Y")
                    rec[f]=clean(v)
                if not rec["Mutation No"] and not rec["Mutation Status"]: continue
                rec["S.No."]=rec["S.No."] or str(len(records)+1)
                for k in ("Mutation Date","Registry Date"): rec[k]=normalize_date(rec[k])
                records.append(rec)
            if not records: raise ValueError("No mutation records were found.")
            self.records=records
            self.source_label.setText(f"●  Imported {Path(path).name}  ·  {len(records)} records")
            self.refresh()
        except Exception as e:
            QMessageBox.critical(self,"Import failed",str(e))

    def paste_data(self):
        dlg=PasteDialog(self)
        if dlg.exec()!=QDialog.DialogCode.Accepted: return
        try:
            data=parse_pasted(dlg.get_text())
            if not data: raise ValueError("No rows detected. Copy cells from Excel and paste again.")
            self.records=data
            self.source_label.setText(f"●  Pasted data  ·  {len(data)} records")
            self.refresh()
        except Exception as e: QMessageBox.warning(self,"Could not read pasted data",str(e))

    def export_csv(self):
        if not self.records:
            QMessageBox.information(self,"Nothing to export","There are no records to export."); return
        path,_=QFileDialog.getSaveFileName(self,"Export filtered report","mutation_pendency.csv","CSV files (*.csv)")
        if not path:return
        try:
            with open(path,"w",newline="",encoding="utf-8-sig") as f:
                writer=csv.DictWriter(f,fieldnames=FIELDS); writer.writeheader()
                for r in self.filtered_records(): writer.writerow(r)
            QMessageBox.information(self,"Export complete",f"Saved {len(self.filtered_records())} rows.")
        except Exception as e: QMessageBox.critical(self,"Export failed",str(e))

    def filtered_records(self):
        text=self.search.text().strip().lower()
        status=self.status_filter.currentText()
        recs=[r for r in self.records if (status=="All statuses" or r.get("Mutation Status","")==status)
              and (not text or any(text in clean(r.get(f,"")).lower() for f in FIELDS))]
        sort=self.sort_filter.currentText()
        if sort=="Latest mutation first":
            recs.sort(key=lambda r: self.date_key(r.get("Mutation Date","")), reverse=True)
        elif sort=="Oldest mutation first":
            recs.sort(key=lambda r: self.date_key(r.get("Mutation Date","")))
        elif sort=="Mutation number ↓":
            recs.sort(key=lambda r:self.number_key(r.get("Mutation No","")),reverse=True)
        else: recs.sort(key=lambda r:self.number_key(r.get("Mutation No","")))
        return recs

    def date_key(self,v):
        try:return datetime.strptime(v,"%d/%m/%Y")
        except:return datetime.min
    def number_key(self,v):
        try:return int(float(v))
        except:return -1

    def refresh(self):
        statuses=sorted({clean(r.get("Mutation Status")) for r in self.records if clean(r.get("Mutation Status"))})
        current=self.status_filter.currentText()
        self.status_filter.blockSignals(True); self.status_filter.clear(); self.status_filter.addItem("All statuses"); self.status_filter.addItems(statuses)
        if current in ["All statuses"]+statuses:self.status_filter.setCurrentText(current)
        self.status_filter.blockSignals(False)
        total=len(self.records)
        count_status=lambda s:sum(1 for r in self.records if r.get("Mutation Status","").lower()==s.lower())
        no_registry=sum(1 for r in self.records if not r.get("Registry Number","").strip() or r.get("Registry Number","").strip()=="-" or not r.get("Registry Date","").strip() or r.get("Registry Date","").strip()=="-")
        self.kpi_total.setText(f"{total:,}")
        self.kpi_incorp.setText(f"{count_status('Incorporated'):,}")
        self.kpi_entry.setText(f"{count_status('Entry Level'):,}")
        self.kpi_registry.setText(f"{no_registry:,}")
        self.model.removeRows(0,self.model.rowCount())
        for rec in self.filtered_records():
            items=[]
            for f in FIELDS:
                item=QStandardItem(clean(rec.get(f,"")))
                item.setEditable(False)
                if f=="Mutation Status":
                    st=clean(rec.get(f,"")).lower()
                    colors={"incorporated":("#e5f8ef","#14875c"),"sanctioned":("#eaf0ff","#356df3"),"entry level":("#fff4df","#a86b0c"),"entry done":("#f0eaff","#7953bd")}
                    bg,fg=colors.get(st,("#f0f3f8","#65738a"))
                    item.setBackground(QColor(bg)); item.setForeground(QColor(fg))
                    font=item.font(); font.setBold(True); item.setFont(font)
                elif f in ("Mutation No","Registry Number"):
                    item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                items.append(item)
            self.model.appendRow(items)
        self.result_count.setText(f"Showing {self.model.rowCount()} of {total} records")
        self.page_info.setText(f"{self.model.rowCount()} rows")
        self.updated_label.setText("Local-first · Data stays on this computer")
        self.table.resizeRowsToContents()

    def apply_filters(self):
        if not hasattr(self,"model"): return
        self.refresh()

def main():
    app=QApplication(sys.argv); app.setStyle("Fusion")
    app.setFont(QFont("Segoe UI",10))
    w=MainWindow(); w.show()
    sys.exit(app.exec())
if __name__=="__main__": main()
